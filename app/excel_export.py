"""Excel exports for report views."""
from __future__ import annotations

import datetime
import re
import time
from copy import copy
from io import BytesIO
from typing import Optional

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from . import config, db

TEMPLATE = config.ASSETS_DIR / "shablon.xlsx"
LEGACY_TEMPLATE = config.BASE_DIR / "shablon.xlsx"
OBSHIY_TEMPLATE = config.ASSETS_DIR / "obshiy_ves_shablon.xlsx"
LEGACY_OBSHIY_TEMPLATE = config.BASE_DIR / "obshiy_ves_shablon.xlsx"
UMUMIY_TEMPLATE = config.ASSETS_DIR / "umumiy_hisobot_shabloni.xlsx"
LEGACY_UMUMIY_TEMPLATE = config.BASE_DIR / "umumiy_hisobot_shabloni.xlsx"
OBSHIY_ACTION_ORDER = ("top", "topchiqgan", "bizda", "chiqgan")

OBSHIY_TYPES = {
    "top", "top'dan chiqgan", "topdan chiqgan", "topchiqgan", "top_dan_chiqgan",
    "bizda qoladigan", "bizda", "bizda_qoladigan", "bizdan chiqgan", "chiqgan", "bizdan_chiqgan",
    "umumiy hisobot"
}


def _safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", (name or "Hisobot")).strip()
    return (name[:31] or "Hisobot")


def _safe_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", (name or "Hisobot")).strip()
    name = re.sub(r"\s+", " ", name)
    return f"{name or 'Hisobot'} KARGOLARGA TARQATISH.xlsx"


def _safe_obshiy_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", (name or "Hisobot")).strip()
    name = re.sub(r"\s+", " ", name)
    return f"{name or 'Hisobot'} OBSHIY VES.xlsx"


def _safe_umumiy_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", (name or "Hisobot")).strip()
    name = re.sub(r"\s+", " ", name)
    return f"{name or 'Hisobot'} UMUMIY HISOBOT.xlsx"


def _num(v) -> float:
    try:
        return round(float(v or 0), 4)
    except (TypeError, ValueError):
        return 0.0


def _formula_num(v) -> str:
    return f"{_num(v):.10g}"


def _entry_expr_and_value(entry: dict) -> tuple[str, float]:
    weight = _num(entry.get("weight"))
    coef = _num(entry.get("coefficient"))
    net = _num(entry.get("net"))
    mode = str(entry.get("coefficient_mode") or "").strip().lower()

    if mode == "box":
        # Box mode: Net is gross, tare is box weight
        return _formula_num(net or weight), net or weight

    if coef > 0:
        if not net:
            net = round(weight - coef, 4)
        return f"{_formula_num(weight)}-{_formula_num(coef)}", net
    return _formula_num(net or weight), net or weight


def _slot_value(slot: dict) -> float:
    return round(float(slot.get("value") or 0) + sum(float(x) for x in slot.get("deltas", [])), 4)


def _slot_cell_value(slot: dict):
    expr = str(slot["expr"])
    deltas = list(slot.get("deltas") or [])
    if not deltas and not slot.get("force_formula"):
        return _num(slot.get("value"))
    parts = [expr]
    for delta in deltas:
        n = _formula_num(abs(delta))
        parts.append(("-" if delta < 0 else "+") + n)
    return "=" + "".join(parts)


def _append_reys(slots: dict[str, list[dict]], entry: dict) -> None:
    tovar_turi = str(entry.get("tovar_turi") or "").strip()
    if not tovar_turi:
        return
    expr, value = _entry_expr_and_value(entry)
    slots.setdefault(tovar_turi, []).append({
        "expr": expr,
        "value": value,
        "deltas": [],
        "force_formula": str(expr) != _formula_num(value),
    })


def _apply_adjust(slots: dict[str, list[dict]], entry: dict) -> None:
    from_type = str(entry.get("from_type") or "").strip()
    to_type = str(entry.get("to_type") or "").strip()
    weight = _num(entry.get("weight"))
    if not from_type or not to_type or weight <= 0:
        return

    slots.setdefault(to_type, []).append({
        "expr": _formula_num(weight),
        "value": weight,
        "deltas": [],
        "force_formula": False,
    })

    remaining = weight
    source = slots.setdefault(from_type, [])
    for slot in reversed(source):
        if remaining <= 0:
            break
        available = max(0.0, _slot_value(slot))
        if available <= 0:
            continue
        take = min(available, remaining)
        slot.setdefault("deltas", []).append(-take)
        slot["force_formula"] = True
        remaining = round(remaining - take, 4)
    if remaining > 0:
        source.append({
            "expr": "0",
            "value": 0.0,
            "deltas": [-remaining],
            "force_formula": True,
        })


def _obshiy_code(entry: dict) -> str:
    return (str(entry.get("tovar_turi") or "").strip() or "Kodsiz")


def _sum_by_code(entries: list[dict]) -> tuple[dict[str, float], list[str]]:
    totals: dict[str, float] = {}
    order: list[str] = []
    for entry in entries:
        code = _obshiy_code(entry)
        if code not in totals:
            totals[code] = 0.0
            order.append(code)
        totals[code] = round(totals[code] + _num(entry.get("weight")), 4)
    return totals, order


def _obshiy_value(entry: dict) -> float:
    return _num(entry.get("weight"))


def _sum_obshiy_values_by_code(entries: list[dict]) -> tuple[dict[str, float], list[str]]:
    totals: dict[str, float] = {}
    order: list[str] = []
    for entry in entries:
        code = _obshiy_code(entry)
        if code not in totals:
            totals[code] = 0.0
            order.append(code)
        totals[code] = round(totals[code] + _obshiy_value(entry), 4)
    return totals, order


def _ordered_codes(*groups: list[str]) -> list[str]:
    seen: set[str] = set()
    order: list[str] = []
    for group in groups:
        for code in group:
            if code not in seen:
                seen.add(code)
                order.append(code)
    return order


def _obshiy_rows(
    base: dict[str, float],
    plus: dict[str, float],
    minus: dict[str, float],
    order: list[str],
) -> list[tuple[str, float, float]]:
    rows: list[tuple[str, float, float]] = []
    for code in order:
        b = base.get(code, 0.0)
        p = plus.get(code, 0.0)
        m = minus.get(code, 0.0)
        transfer = round(p - m, 4)
        if b != 0 or transfer != 0:
            rows.append((code, b, transfer))
    return rows


def _write_obshiy_sheet(ws, rows: list[tuple[str, float, float]], context: dict[str, str]) -> None:
    ws.cell(1, 1).value = "kod"
    ws.cell(1, 2).value = context["sheet"]
    ws.cell(1, 3).value = context["column"]
    ws.cell(1, 4).value = "jami"

    for r in range(2, max(ws.max_row, len(rows) + 2) + 1):
        for c in range(1, 5):
            ws.cell(r, c).value = None

    for offset, (code, base, transfer) in enumerate(rows):
        idx = 2 + offset
        ws.cell(idx, 1).value = code
        ws.cell(idx, 2).value = base
        ws.cell(idx, 3).value = transfer
        ws.cell(idx, 4).value = None
        for col in (2, 3):
            ws.cell(idx, col).number_format = "0.00"

    total_row = 2
    last_row = max(total_row, len(rows) + 1)
    total_cell = ws.cell(total_row, 4)
    total_cell.value = f"=SUM(B{total_row}:C{last_row})" if rows else 0
    total_cell.number_format = "0.00"


def render_obshiy_excel(report_name: str, entries_by_action: dict[str, list[dict]]) -> tuple[bytes, str]:
    from openpyxl import Workbook, load_workbook

    top, top_order = _sum_obshiy_values_by_code(entries_by_action.get("top", []))
    topchiqgan, topchiqgan_order = _sum_obshiy_values_by_code(entries_by_action.get("topchiqgan", []))
    bizda, bizda_order = _sum_obshiy_values_by_code(entries_by_action.get("bizda", []))
    chiqgan, chiqgan_order = _sum_obshiy_values_by_code(entries_by_action.get("chiqgan", []))

    top_rows = _obshiy_rows(
        top,
        plus=chiqgan,
        minus=topchiqgan,
        order=_ordered_codes(top_order, chiqgan_order, topchiqgan_order),
    )
    bizda_rows = _obshiy_rows(
        bizda,
        plus=topchiqgan,
        minus=chiqgan,
        order=_ordered_codes(bizda_order, topchiqgan_order, chiqgan_order),
    )

    template = OBSHIY_TEMPLATE if OBSHIY_TEMPLATE.exists() else LEGACY_OBSHIY_TEMPLATE
    if template.exists():
        wb = load_workbook(template)
    else:
        wb = Workbook()
        wb.active.title = "top"
        wb.create_sheet("bizda qoladigan")

    while len(wb.worksheets) < 2:
        wb.create_sheet("bizda qoladigan" if len(wb.worksheets) == 1 else f"Hisobot {len(wb.worksheets) + 1}")

    _write_obshiy_sheet(
        wb.worksheets[0],
        top_rows,
        {"sheet": "top", "column": "bizdan chiqgan"},
    )
    _write_obshiy_sheet(
        wb.worksheets[1],
        bizda_rows,
        {"sheet": "bizda qoladigan", "column": "topdan chiqgan"},
    )
    for ws in wb.worksheets[2:]:
        wb.remove(ws)

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), _safe_obshiy_filename(report_name)


def build_obshiy_excel(report_id: int) -> tuple[bytes, str]:
    report_name = db.report_name(report_id) or "Hisobot"
    entries_by_action = {
        action: list(reversed(db.list_entries(report_id, action, limit=2000)))
        for action in OBSHIY_ACTION_ORDER
    }
    return render_obshiy_excel(report_name, entries_by_action)


async def build_obshiy_excel_async(session: AsyncSession, reys_id: int | None = None) -> tuple[bytes, str]:
    from .models.cargo import Cargo
    from .models.entry import Entry
    from .models.reys import Reys

    target_id = reys_id
    if not target_id:
        latest = (await session.execute(
            select(Reys).where(Reys.deleted_at.is_(None)).order_by(Reys.id.desc()).limit(1)
        )).scalar_one_or_none()
        if latest:
            target_id = latest.id
        else:
            return render_obshiy_excel("Hisobot", {a: [] for a in OBSHIY_ACTION_ORDER})

    reys = (await session.execute(
        select(Reys).where(Reys.id == target_id, Reys.deleted_at.is_(None))
    )).scalar_one_or_none()
    if not reys:
        return build_obshiy_excel(target_id)

    cargo_code = ""
    if reys.cargo_id:
        cargo = (await session.execute(select(Cargo).where(Cargo.id == reys.cargo_id))).scalar_one_or_none()
        if cargo:
            cargo_code = cargo.code
    report_name = f"{cargo_code + ' ' if cargo_code else ''}{reys.code}{' ' + reys.custom_name if reys.custom_name else ''}".strip()

    entries_models = (await session.execute(
        select(Entry).where(Entry.reys_id == reys.id, Entry.deleted_at.is_(None)).order_by(Entry.id.asc())
    )).scalars().all()

    entries_by_action: dict[str, list[dict]] = {a: [] for a in OBSHIY_ACTION_ORDER}
    for e in entries_models:
        t = (e.tovar_turi or "").strip().lower()
        cat = None
        if t == "top":
            cat = "top"
        elif t in ("top'dan chiqgan", "topdan chiqgan", "topchiqgan", "top_dan_chiqgan"):
            cat = "topchiqgan"
        elif t in ("bizda qoladigan", "bizda", "bizda_qoladigan"):
            cat = "bizda"
        elif t in ("bizdan chiqgan", "chiqgan", "bizdan_chiqgan"):
            cat = "chiqgan"
        if cat:
            entries_by_action[cat].append({
                "tovar_turi": e.box_code or "Kodsiz",
                "weight": e.gross_weight,
                "coefficient": e.tare_weight,
                "box_weight": e.tare_weight,
                "net": e.net_weight,
            })

    return render_obshiy_excel(report_name, entries_by_action)


def _summary_type_key(tovar_turi: str) -> str:
    key = str(tovar_turi or "").strip().lower()
    if key == "one":
        return "oneway"
    if key == "uztez":
        return "uzt"
    if key.startswith("x"):
        return "xabib"
    return key


def _inventory_for_summary(report_id: int) -> dict[str, float]:
    inv: dict[str, float] = {}
    for tovar_turi, value in db.get_inventory(report_id).items():
        key = _summary_type_key(tovar_turi)
        if not key:
            continue
        inv[key] = round(inv.get(key, 0) + _num(value), 4)
    return inv


def _summary_box_weight(obshiy_entries: dict[str, list[dict]], reys_entries: list[dict]) -> float:
    obshiy_box = round(
        sum(_num(e.get("box_weight")) or _num(e.get("coefficient")) for e in obshiy_entries.get("bizda", [])),
        4,
    )
    kargo_box = round(sum(_num(e.get("box_weight")) for e in reys_entries), 4)
    return round(obshiy_box - kargo_box, 4)


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.fill = copy(src.fill)
        dst.font = copy(src.font)
        dst.border = copy(src.border)


def render_umumiy_excel(
    report_name: str,
    obshiy_entries: dict[str, list[dict]],
    reys_entries: list[dict],
    inv: dict[str, float],
) -> tuple[bytes, str]:
    from datetime import date
    from openpyxl import Workbook, load_workbook

    topchiqgan, topchiqgan_order = _sum_obshiy_values_by_code(obshiy_entries.get("topchiqgan", []))
    bizda, bizda_order = _sum_obshiy_values_by_code(obshiy_entries.get("bizda", []))
    chiqgan, chiqgan_order = _sum_obshiy_values_by_code(obshiy_entries.get("chiqgan", []))
    bizda_rows = _obshiy_rows(
        bizda,
        plus=topchiqgan,
        minus=chiqgan,
        order=_ordered_codes(bizda_order, topchiqgan_order, chiqgan_order),
    )

    bizda_total = round(sum(round(base + transfer, 4) for _, base, transfer in bizda_rows), 4)
    box_weight_total = _summary_box_weight(obshiy_entries, reys_entries)
    normalized_inv = {_summary_type_key(k): _num(v) for k, v in inv.items() if _summary_type_key(k)}
    top_inventory = _num(normalized_inv.get("top", 0))
    normalized_inv["top"] = 0.0

    template = UMUMIY_TEMPLATE if UMUMIY_TEMPLATE.exists() else LEGACY_UMUMIY_TEMPLATE
    if template.exists():
        wb = load_workbook(template)
    else:
        wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name(report_name)

    max_clear_row = max(ws.max_row, 129)
    for r in range(2, max_clear_row + 1):
        ws.cell(r, 1).value = None

    a_row = 2
    clean_total = round(bizda_total - top_inventory, 4)
    ws.cell(a_row, 1).value = clean_total
    ws.cell(a_row, 1).number_format = "0.00"

    ws["E2"] = "To'lashi kerak bo'lgan summa:"
    ws["F2"] = date.today().strftime("%d.%m.%Y")
    ws["F3"] = report_name
    ws["C2"] = "=SUM(A:A)"
    ws["C3"] = box_weight_total
    ws["C3"].number_format = "0.00"

    label_rows: dict[str, int] = {}
    for row in range(4, ws.max_row + 1):
        raw_label = str(ws.cell(row, 2).value or "").strip().lower()
        label = _summary_type_key(raw_label)
        if label:
            if raw_label == "one":
                ws.cell(row, 2).value = "oneway"
            label_rows[label] = row

    represented = set(label_rows)
    custom_types = [
        t for t, value in normalized_inv.items()
        if value and t not in represented and t not in {"mandarin", "uztez"}
    ]
    next_row = max([r for r in label_rows.values()] + [15]) + 1
    for tovar_turi in custom_types:
        _copy_row_style(ws, 15, next_row, 8)
        ws.cell(next_row, 2).value = tovar_turi
        label_rows[tovar_turi] = next_row
        next_row += 1

    distributed_labels = {"akb", "jet", "xabib", "navo", "jon", "oneway", "redwing", "uzt"}
    for label, row in label_rows.items():
        if label in ("karobka", "mandarin"):
            continue
        ws.cell(row, 3).value = normalized_inv.get(label, 0)
        ws.cell(row, 3).number_format = "0.00"

    distributable_rows = [
        row for label, row in label_rows.items()
        if label in distributed_labels
    ]
    non_distributed_rows = [
        row for label, row in label_rows.items()
        if label not in distributed_labels and label not in {"karobka", "mandarin"}
    ]

    distributed_formula = "+".join(f"C{r}" for r in sorted(distributable_rows)) or "0"
    non_distributed_formula = "+".join(f"C{r}" for r in sorted(non_distributed_rows)) or "0"

    ws["C13"] = f"={distributed_formula}"
    ws["C13"].number_format = "0.00"
    ws["C14"] = f"={non_distributed_formula}"
    ws["C14"].number_format = "0.00"

    # mandarin = total - distributed - non_distributed
    ws["C12"] = "=(A2-C13)-C14"
    ws["C12"].number_format = "0.00"

    ws.sheet_view.showGridLines = True
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), _safe_umumiy_filename(report_name)


def render_multi_reys_summary(reyslar: list, cargos: dict[int, str]) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Barcha Reyslar"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="FDE047")
    hdr_font = Font(bold=True, size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    headers = [
        "T/r", "Reys Kodi", "Kargo", "Qo'shimcha Nomi",
        "Sana", "Sof Vazn (kg)", "Karobka Vazni (kg)",
        "Jami Og'irlik (kg)", "Karobkalar Soni"
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = align_center
        cell.border = border

    row_idx = 2
    for idx, r in enumerate(reyslar, start=1):
        cargo_name = cargos.get(r.cargo_id, "-") if r.cargo_id else "-"
        toza = float(r.toza_kg or 0.0)
        karobka = float(r.karobka_plus_kg or 0.0)
        jami = round(toza + karobka, 2)
        boxes = len(r.entries) if getattr(r, "entries", None) else 0

        vals = [
            (idx, align_center, None),
            (r.code, align_center, None),
            (cargo_name, align_center, None),
            (r.custom_name or "-", align_center, None),
            (str(r.date or "-"), align_center, None),
            (toza, align_right, "0.00"),
            (karobka, align_right, "0.00"),
            (jami, align_right, "0.00"),
            (boxes, align_right, "#,##0"),
        ]
        for col_idx, (val, alignment, num_fmt) in enumerate(vals, start=1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.alignment = alignment
            cell.border = border
            if num_fmt:
                cell.number_format = num_fmt
        row_idx += 1

    if len(reyslar) > 0:
        last_data_row = row_idx - 1
        ws.cell(row_idx, 1, "").border = border
        ws.cell(row_idx, 2, "JAMI:").border = border
        ws.cell(row_idx, 2).font = Font(bold=True)
        for c in range(3, 6):
            ws.cell(row_idx, c, "").border = border
        for c_idx, letter in [(6, "F"), (7, "G"), (8, "H"), (9, "I")]:
            cell = ws.cell(row_idx, c_idx)
            cell.value = f"=SUM({letter}2:{letter}{last_data_row})"
            cell.font = Font(bold=True)
            cell.alignment = align_right
            cell.border = border
            cell.number_format = "0.00" if c_idx < 9 else "#,##0"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.sheet_view.showGridLines = True
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), "BARCHA REYSLAR UMUMIY HISOBOT.xlsx"


def build_umumiy_excel(report_id: int) -> tuple[bytes, str]:
    report_name = db.report_name(report_id) or "Hisobot"
    obshiy_entries = {
        action: list(reversed(db.list_entries(report_id, action, limit=2000)))
        for action in OBSHIY_ACTION_ORDER
    }
    reys_entries = list(reversed(db.list_entries(report_id, "reys", limit=2000)))
    inv = _inventory_for_summary(report_id)
    return render_umumiy_excel(report_name, obshiy_entries, reys_entries, inv)


async def build_umumiy_excel_async(session: AsyncSession, reys_id: int | None = None) -> tuple[bytes, str]:
    from .models.cargo import Cargo
    from .models.entry import Entry
    from .models.inventory import Inventory
    from .models.reys import Reys

    if reys_id:
        reys = (await session.execute(
            select(Reys).where(Reys.id == reys_id, Reys.deleted_at.is_(None))
        )).scalar_one_or_none()
        if not reys:
            return build_umumiy_excel(reys_id)

        cargo_code = ""
        if reys.cargo_id:
            cargo = (await session.execute(select(Cargo).where(Cargo.id == reys.cargo_id))).scalar_one_or_none()
            if cargo:
                cargo_code = cargo.code
        report_name = f"{cargo_code + ' ' if cargo_code else ''}{reys.code}{' ' + reys.custom_name if reys.custom_name else ''}".strip()

        entries_models = (await session.execute(
            select(Entry).where(Entry.reys_id == reys.id, Entry.deleted_at.is_(None)).order_by(Entry.id.asc())
        )).scalars().all()

        obshiy_entries: dict[str, list[dict]] = {a: [] for a in OBSHIY_ACTION_ORDER}
        reys_entries: list[dict] = []
        for e in entries_models:
            t = (e.tovar_turi or "").strip().lower()
            cat = None
            if t == "top":
                cat = "top"
            elif t in ("top'dan chiqgan", "topdan chiqgan", "topchiqgan", "top_dan_chiqgan"):
                cat = "topchiqgan"
            elif t in ("bizda qoladigan", "bizda", "bizda_qoladigan"):
                cat = "bizda"
            elif t in ("bizdan chiqgan", "chiqgan", "bizdan_chiqgan"):
                cat = "chiqgan"
            if cat:
                obshiy_entries[cat].append({
                    "tovar_turi": e.box_code or "Kodsiz",
                    "weight": e.gross_weight,
                    "coefficient": e.tare_weight,
                    "box_weight": e.tare_weight,
                    "net": e.net_weight,
                })
            else:
                reys_entries.append({
                    "tovar_turi": e.tovar_turi,
                    "weight": e.gross_weight,
                    "coefficient": e.tare_weight,
                    "box_weight": e.tare_weight,
                    "net": e.net_weight,
                })

        inv_models = (await session.execute(
            select(Inventory).where(Inventory.reys_id == reys.id)
        )).scalars().all()
        inv = {i.tovar_turi: float(i.weight) for i in inv_models}

        return render_umumiy_excel(report_name, obshiy_entries, reys_entries, inv)
    else:
        reyslar = (await session.execute(
            select(Reys).where(Reys.deleted_at.is_(None)).order_by(Reys.id.desc())
        )).scalars().all()
        if not reyslar:
            return render_umumiy_excel("Hisobot", {a: [] for a in OBSHIY_ACTION_ORDER}, [], {})

        cargos = {c.id: c.code for c in (await session.execute(select(Cargo))).scalars().all()}
        return render_multi_reys_summary(reyslar, cargos)


def render_kargo_excel(report_name: str, entries: list[dict], adjusts: list[dict]) -> tuple[bytes, str]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    slots: dict[str, list[dict]] = {}
    custom_types: list[str] = []
    default_set = set(db.DEFAULT_TYPES)
    for entry in entries:
        tovar_turi = str(entry.get("tovar_turi") or "").strip()
        if not tovar_turi:
            continue
        _append_reys(slots, entry)
        if tovar_turi not in default_set and tovar_turi not in custom_types:
            custom_types.append(tovar_turi)
    for entry in adjusts:
        _apply_adjust(slots, entry)
        for tovar_turi in (str(entry.get("from_type") or "").strip(), str(entry.get("to_type") or "").strip()):
            if tovar_turi and tovar_turi not in default_set and tovar_turi not in custom_types:
                custom_types.append(tovar_turi)

    template = TEMPLATE if TEMPLATE.exists() else LEGACY_TEMPLATE
    if template.exists():
        wb = load_workbook(template)
    else:
        wb = Workbook()
        ws0 = wb.active
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for r, tovar_turi in enumerate(db.DEFAULT_TYPES, start=1):
            ws0.cell(r, 1).value = tovar_turi
            ws0.cell(r, 1).fill = PatternFill("solid", fgColor="FFFF00")
            ws0.cell(r, 1).font = Font(bold=True)
            ws0.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws0.cell(r, 1).border = border
            ws0.cell(r, 2).number_format = "0.00"
            ws0.cell(r, 2).border = border
        ws0.column_dimensions["A"].width = 18
        for col in range(2, 30):
            ws0.column_dimensions[get_column_letter(col)].width = 12

    ws = wb.active
    ws.title = _safe_sheet_name(report_name)

    label_style = copy(ws["A1"]._style)
    custom_label_style = copy(ws["A13"]._style if ws["A13"].has_style else ws["A1"]._style)
    value_style = copy(ws["B1"]._style)
    blank_style = copy(ws["K1"]._style)
    label_fill = copy(ws["A1"].fill)
    value_num_format = ws["B1"].number_format

    types = list(db.DEFAULT_TYPES)
    data_rows: list[tuple[int, str, bool]] = []
    row = 1
    for tovar_turi in types:
        data_rows.append((row, tovar_turi, False))
        row += 1
    if custom_types:
        row += 1
        for tovar_turi in custom_types:
            data_rows.append((row, tovar_turi, True))
            row += 1

    max_entries = max([len(slots.get(t, [])) for _, t, _ in data_rows] + [1])
    max_col = max(2, 1 + max_entries)
    total_start = (data_rows[-1][0] if data_rows else 1) + 6
    total_end = total_start + len(data_rows) - 1
    clear_rows = max(ws.max_row, total_end)
    clear_cols = max(ws.max_column, max_col)

    for clear_row in range(1, clear_rows + 1):
        for clear_col in range(1, clear_cols + 1):
            cell = ws.cell(clear_row, clear_col)
            cell.value = None
            cell._style = copy(blank_style)

    data_row_by_type: dict[str, int] = {}
    for data_row, tovar_turi, is_custom in data_rows:
        data_row_by_type[tovar_turi] = data_row
        cell = ws.cell(data_row, 1)
        cell.value = tovar_turi
        cell._style = copy(custom_label_style if is_custom else label_style)
        cell.fill = copy(label_fill)
        for idx, slot in enumerate(slots.get(tovar_turi, []), start=2):
            val_cell = ws.cell(data_row, idx)
            val_cell.value = _slot_cell_value(slot)
            val_cell._style = copy(value_style)
            val_cell.number_format = value_num_format

    total_row = total_start
    last_col = get_column_letter(max_col)
    for _, tovar_turi, is_custom in data_rows:
        a = ws.cell(total_row, 1)
        b = ws.cell(total_row, 2)
        a.value = tovar_turi
        a._style = copy(custom_label_style if is_custom else label_style)
        a.fill = copy(label_fill)
        b.value = f"=SUM(B{data_row_by_type[tovar_turi]}:{last_col}{data_row_by_type[tovar_turi]})"
        b._style = copy(value_style)
        b.number_format = value_num_format
        total_row += 1

    ws.sheet_view.showGridLines = True
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), _safe_filename(report_name)


def build_kargo_excel(report_id: int) -> tuple[bytes, str]:
    report_name = db.report_name(report_id) or "Hisobot"
    entries = list(reversed(db.list_entries(report_id, "reys", limit=2000)))
    adjusts = list(reversed(db.list_entries(report_id, "adjust", limit=2000)))
    return render_kargo_excel(report_name, entries, adjusts)


async def build_kargo_excel_async(
    session: AsyncSession,
    reys_id: int | None = None,
    start: str | None = None,
    end: str | None = None,
) -> tuple[bytes, str]:
    from .models.cargo import Cargo
    from .models.entry import Entry
    from .models.reys import Reys
    from .models.activity import ActivityLog

    if start and end and not reys_id:
        try:
            start_ts = int(datetime.datetime.strptime(start, "%Y-%m-%d").replace(hour=0, minute=0, second=0).timestamp())
            end_ts = int(datetime.datetime.strptime(end, "%Y-%m-%d").replace(hour=23, minute=59, second=59).timestamp())
        except Exception:
            start_ts, end_ts = 0, int(time.time())
        report_name = f"KARGOLAR {start} - {end}"
        entries_models = (await session.execute(
            select(Entry)
            .where(Entry.created_at >= start_ts, Entry.created_at <= end_ts, Entry.deleted_at.is_(None))
            .order_by(Entry.id.asc())
        )).scalars().all()
        adjusts_models = []
    else:
        target_id = reys_id
        if not target_id:
            latest = (await session.execute(
                select(Reys).where(Reys.deleted_at.is_(None)).order_by(Reys.id.desc()).limit(1)
            )).scalar_one_or_none()
            if latest:
                target_id = latest.id
            else:
                return render_kargo_excel("Hisobot", [], [])

        reys = (await session.execute(
            select(Reys).where(Reys.id == target_id, Reys.deleted_at.is_(None))
        )).scalar_one_or_none()
        if not reys:
            return build_kargo_excel(target_id)

        cargo_code = ""
        if reys.cargo_id:
            cargo = (await session.execute(select(Cargo).where(Cargo.id == reys.cargo_id))).scalar_one_or_none()
            if cargo:
                cargo_code = cargo.code
        report_name = f"{cargo_code + ' ' if cargo_code else ''}{reys.code}{' ' + reys.custom_name if reys.custom_name else ''}".strip()

        entries_models = (await session.execute(
            select(Entry).where(Entry.reys_id == reys.id, Entry.deleted_at.is_(None)).order_by(Entry.id.asc())
        )).scalars().all()
        adjusts_models = (await session.execute(
            select(ActivityLog).where(ActivityLog.reys_id == reys.id, ActivityLog.action == "adjust").order_by(ActivityLog.id.asc())
        )).scalars().all()

    entries = [
        {
            "tovar_turi": e.tovar_turi,
            "weight": e.gross_weight,
            "coefficient": e.tare_weight,
            "net": e.net_weight,
            "box_weight": e.tare_weight,
            "coefficient_mode": e.coefficient_mode,
        }
        for e in entries_models
        if (e.tovar_turi or "").strip().lower() not in OBSHIY_TYPES
    ]
    adjusts = [
        {
            "from_type": a.from_type,
            "to_type": a.to_type,
            "weight": a.weight,
        }
        for a in adjusts_models
        if a.from_type and a.to_type
    ]
    return render_kargo_excel(report_name, entries, adjusts)
