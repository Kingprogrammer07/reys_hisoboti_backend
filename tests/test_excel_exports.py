from __future__ import annotations

from datetime import date
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook


async def _seed_excel_export_data(client: httpx.AsyncClient, auth_headers: dict[str, str]) -> int:
    cargo_resp = await client.post(
        "/api/cargos",
        json={"code": "EXCEL-CARGO-01"},
        headers=auth_headers,
    )
    assert cargo_resp.status_code == 201
    cargo_id = cargo_resp.json()["id"]

    reys_resp = await client.post(
        "/api/reys",
        json={
            "cargo_id": cargo_id,
            "code": "EXCEL-REYS-01",
            "date": "2026-09-25",
            "custom_name": "Excel Smoke",
        },
        headers=auth_headers,
    )
    assert reys_resp.status_code == 201
    reys_id = reys_resp.json()["id"]

    entries = [
        ("BOX-KARGO-01", "akb", 30.0, 1.22),
        ("BOX-KARGO-02", "mandarin", 22.5, 0.94),
        ("BOX-TOP-01", "top", 18.0, 1.0),
        ("BOX-BIZDA-01", "bizda_qoladigan", 12.0, 0.5),
        ("BOX-CHIQGAN-01", "bizdan_chiqgan", 7.0, 0.5),
        ("BOX-TOPCHIQ-01", "top_dan_chiqgan", 3.0, 0.2),
    ]
    for box_code, tovar_turi, gross_weight, tare_weight in entries:
        resp = await client.post(
            "/api/entries/json",
            json={
                "reys_id": reys_id,
                "box_code": box_code,
                "tovar_turi": tovar_turi,
                "gross_weight": gross_weight,
                "tare_weight": tare_weight,
                "coefficient_mode": "box",
                "created_by": "pytest",
            },
            headers=auth_headers,
        )
        assert resp.status_code == 201

    adjust_resp = await client.post(
        "/api/entries/adjust",
        json={
            "reys_id": reys_id,
            "from_type": "akb",
            "to_type": "mandarin",
            "weight": 5.0,
            "created_by": "pytest",
        },
        headers=auth_headers,
    )
    assert adjust_resp.status_code == 200

    return reys_id


def _assert_valid_xlsx(content: bytes, expected_fragment: str) -> None:
    assert content.startswith(b"PK"), "xlsx should be a zipped Office document"

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames
    assert any(
        cell.value is not None
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    ), f"{expected_fragment} workbook should contain data"


@pytest.mark.asyncio
async def test_excel_export_endpoints_return_valid_workbooks(
    client: httpx.AsyncClient,
    auth_headers: dict[str, str],
):
    reys_id = await _seed_excel_export_data(client, auth_headers)
    today = date.today().isoformat()

    endpoints = [
        (f"/api/export/kargo?report_id={reys_id}", "KARGOLARGA"),
        (f"/api/export/obshiy?report_id={reys_id}", "OBSHIY"),
        (f"/api/export/summary?report_id={reys_id}", "UMUMIY"),
        ("/api/export/summary", "BARCHA"),
        (f"/api/export/kargo?start={today}&end={today}", "KARGOLAR"),
    ]

    for path, filename_fragment in endpoints:
        resp = await client.get(path, headers=auth_headers)
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        content_disposition = resp.headers.get("content-disposition", "")
        assert ".xlsx" in content_disposition
        assert filename_fragment.lower() in content_disposition.lower()
        _assert_valid_xlsx(resp.content, filename_fragment)
